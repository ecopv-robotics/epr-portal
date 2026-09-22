import base64
from io import BytesIO

from docx import Document


def login(client):
    response = client.post("/api/auth/login", json={"account": "demo@eprflow.com", "password": "Demo123!"})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def create_application(client, headers):
    response = client.post(
        "/api/applications",
        json={"country": "Germany", "service_type": "WEEE"},
        headers=headers,
    )
    assert response.status_code == 201
    return response.json()


def test_weee_authorization_reports_missing_fields(client):
    headers = login(client)
    application = create_application(client, headers)

    response = client.get(
        f"/api/applications/{application['id']}/authorization/germany-weee",
        headers=headers,
    )

    assert response.status_code == 422
    assert "公司英文名称" in response.json()["detail"]["missing_fields"]


def test_weee_authorization_fills_customer_slots_and_removes_helper_text(client):
    headers = login(client)
    application = create_application(client, headers)
    one_pixel_png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    signature = client.post(
        "/api/files/upload",
        headers=headers,
        data={"purpose": "signature"},
        files={"file": ("signature.png", one_pixel_png, "image/png")},
    )
    assert signature.status_code == 200

    company = client.put(
        f"/api/applications/{application['id']}/company",
        headers=headers,
        json={
            "company_name": "示例科技有限公司",
            "company_name_en": "Example Technology Co., Ltd.",
            "registration_region": "MAINLAND_CHINA",
            "company_type": "有限责任公司",
            "register_number": "91310000TEST123456",
            "established_date": "2021-05-06",
            "register_address": "上海市浦东新区示例路 88 号",
        },
    )
    assert company.status_code == 200
    legal = client.put(
        f"/api/applications/{application['id']}/legal-person",
        headers=headers,
        json={
            "certificate_type": "PASSPORT",
            "gender": "MALE",
            "name_cn": "张三",
            "name_en": "San Zhang",
            "contact_email": "contact@example.com",
            "contact_phone": "+86 138 0000 0000",
            "front_file": "/files/passport.jpg",
            "signature_file": signature.json()["file_url"],
        },
    )
    assert legal.status_code == 200
    store = client.put(
        f"/api/applications/{application['id']}/store",
        headers=headers,
        json={
            "platform": "Amazon",
            "store_url": "https://example.com/store",
            "company_name_en": "Example Technology Co., Ltd.",
            "register_city_en": "Shanghai",
            "postal_code": "200120",
            "register_address_en": "No. 88 Example Road, Pudong New Area",
            "screenshot_file": "/files/store.png",
        },
    )
    assert store.status_code == 200

    response = client.get(
        f"/api/applications/{application['id']}/authorization/germany-weee",
        headers=headers,
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = Document(BytesIO(response.content))
    values = [row.cells[1].text for row in document.tables[0].rows]
    assert values == [
        "Example Technology Co., Ltd.",
        "No. 88 Example Road, Pudong New Area",
        "200120 Shanghai",
        "+86 138 0000 0000",
        "contact@example.com",
        "San Zhang",
        "91310000TEST123456",
    ]
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "城市 拼音" not in body_text
    assert "日期手写" not in body_text
    assert "法人签字" not in body_text
    assert "公司英文名" not in body_text
    assert "Example Technology Co., Ltd." in body_text

