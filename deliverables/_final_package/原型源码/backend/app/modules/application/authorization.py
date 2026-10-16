from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

from app.core.config import get_settings
from app.models import Application


TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "assets" / "germany_weee_authorization_non_german.docx"


class AuthorizationValidationError(ValueError):
    def __init__(self, missing_fields: list[str]):
        super().__init__("授权书资料不完整")
        self.missing_fields = missing_fields


def _clean(value: str | None) -> str:
    return " ".join((value or "").split())


def _required_authorization_values(application: Application) -> dict[str, str]:
    company = application.company
    store = application.store
    legal = company.legal_person if company else None
    values = {
        "company_name": _clean((store.company_name_en if store else None) or (company.company_name_en if company else None)),
        "address": _clean(store.register_address_en if store else None),
        "postal_city": _clean(" ".join(filter(None, [store.postal_code if store else None, store.register_city_en if store else None]))),
        "city": _clean(store.register_city_en if store else None),
        "phone": _clean(legal.contact_phone if legal else None),
        "email": _clean(legal.contact_email if legal else None),
        "legal_name": _clean(legal.name_en if legal else None),
        "registration_number": _clean(company.register_number if company else None),
        "signature_file": _clean(legal.signature_file if legal else None),
    }
    labels = {
        "company_name": "公司英文名称",
        "address": "英文注册地址",
        "postal_city": "邮编和英文城市",
        "phone": "公司联系人电话",
        "email": "公司联系人邮箱",
        "legal_name": "法人拼音名",
        "registration_number": "营业执照号 / 注册号",
        "signature_file": "法人代表签名",
    }
    missing = [label for key, label in labels.items() if not values[key]]
    if application.country != "Germany" or application.service_type != "WEEE":
        missing.insert(0, "仅德国 WEEE 申请支持生成此授权书")
    if missing:
        raise AuthorizationValidationError(missing)
    return values


def _resolve_uploaded_file(file_url: str) -> Path:
    settings = get_settings()
    storage_root = Path(settings.storage_path)
    if not storage_root.is_absolute():
        storage_root = Path.cwd() / storage_root
    storage_root = storage_root.resolve()
    relative = file_url.split("?", 1)[0].removeprefix("/files/").lstrip("/")
    candidate = (storage_root / relative).resolve()
    if storage_root not in candidate.parents or not candidate.is_file():
        raise AuthorizationValidationError(["法人签名文件无法读取，请重新上传"])
    return candidate


def _set_cell_value(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    first_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    first_run.text = value
    first_run.font.color.rgb = RGBColor(0, 0, 0)
    for run in paragraph.runs[1:]:
        run.text = ""
    for extra_paragraph in cell.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _is_red(run) -> bool:
    color = run.font.color.rgb
    return color is not None and str(color).upper() == "FF0000"


def _black(run) -> None:
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "Arial"
    run.font.size = Pt(10)


def generate_germany_weee_authorization(application: Application, generated_at: datetime | None = None) -> BytesIO:
    values = _required_authorization_values(application)
    signature_path = _resolve_uploaded_file(values["signature_file"])
    if not TEMPLATE_PATH.is_file():
        raise RuntimeError("德国 WEEE 授权书模板不存在")

    document = Document(TEMPLATE_PATH)
    customer_table = document.tables[0]
    customer_values = [
        values["company_name"], values["address"], values["postal_city"], values["phone"],
        values["email"], values["legal_name"], values["registration_number"],
    ]
    for row, value in zip(customer_table.rows, customer_values, strict=True):
        _set_cell_value(row.cells[1], value)

    date_text = (generated_at or datetime.now()).strftime("%d.%m.%Y")
    for paragraph in document.paragraphs:
        original_text = paragraph.text

        if "城市 拼音" in original_text and "日期手写" in original_text:
            paragraph.clear()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(f"{values['city']}, {date_text}")
            run.underline = True
            _black(run)
            continue

        if "法人签字" in original_text and "CEO" in original_text:
            paragraph.clear()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.add_run().add_picture(str(signature_path), height=Inches(0.42))
            role = paragraph.add_run(f"  {values['legal_name']}, CEO")
            _black(role)
            continue

        if "ECOPV-EU GmbH" in original_text and "公司英文名" in original_text:
            for run in paragraph.runs:
                if _is_red(run):
                    run.text = values["company_name"] if "公司英文名" in run.text else ""
                    _black(run)
            continue

        if original_text.count("Company name") >= 2:
            red_runs = [run for run in paragraph.runs if _is_red(run)]
            for index, run in enumerate(red_runs):
                run.text = "Company name" if index == 0 else ""
                _black(run)
            continue

        for run in paragraph.runs:
            if _is_red(run):
                run.text = ""

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output

