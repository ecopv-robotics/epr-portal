from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\vibecoding-project\epr-portal")
PACKAGE = ROOT / "deliverables" / "EPR下单系统业务流程及原型_2026-10-12"
DIAGRAMS = PACKAGE / "业务流程图"
SCREENSHOTS = PACKAGE / "原型截图"
ASSETS = PACKAGE / "文档资源"
DOCX_PATH = PACKAGE / "EPR下单系统业务流程与原型确认稿.docx"

BLUE = "0B63CE"
NAVY = "123A69"
LIGHT_BLUE = "EAF3FF"
PALE_BLUE = "F5F9FF"
GRAY = "666666"
BORDER = "D9D9D9"
ROLE_COLORS = {
    "客户": "DCEBFF",
    "销售": "FFF0D9",
    "系统管理员": "EEE6FF",
    "财务": "DCF7EC",
    "商务": "E4F1FF",
    "交付": "FFE4E1",
}


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_centered(draw: ImageDraw.ImageDraw, box, text: str, text_font, fill="#1F2937"):
    x1, y1, x2, y2 = box
    lines = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=11) or [""])
    line_height = int(text_font.size * 1.35)
    total = len(lines) * line_height
    y = y1 + (y2 - y1 - total) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=text_font)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=text_font, fill=fill)
        y += line_height


def arrow(draw: ImageDraw.ImageDraw, start, end, color="#5B7898", width=5):
    x1, y1 = start
    x2, y2 = end
    if abs(y2 - y1) < 8:
        points = [(x1, y1), (x2, y2)]
    else:
        mid_x = (x1 + x2) / 2
        points = [(x1, y1), (mid_x, y1), (mid_x, y2), (x2, y2)]
    draw.line(points, fill=color, width=width, joint="curve")
    head = 14
    draw.polygon([(x2, y2), (x2 - head, y2 - 9), (x2 - head, y2 + 9)], fill=color)


def make_overall_flow(path: Path):
    width, height = 2300, 1160
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font, lane_font, node_font = font(46, True), font(30, True), font(27, False)
    draw.text((70, 35), "EPR 下单系统端到端业务流程", font=title_font, fill="#102A43")
    roles = ["客户", "销售", "系统管理员", "财务", "商务", "交付"]
    lane_top, lane_h = 120, 165
    for i, role in enumerate(roles):
        y1 = lane_top + i * lane_h
        y2 = y1 + lane_h
        draw.rectangle((40, y1, width - 40, y2), fill="#FAFCFF" if i % 2 == 0 else "#F4F7FB", outline="#D8E1EB", width=2)
        draw.rounded_rectangle((55, y1 + 35, 220, y2 - 35), radius=18, fill="#" + ROLE_COLORS[role], outline="#A8B7C7", width=2)
        draw_centered(draw, (55, y1 + 35, 220, y2 - 35), role, lane_font)

    nodes = [
        ("客户", 285, "选购服务\n加入购物车"),
        ("客户", 520, "提交订单\n默认待付款"),
        ("销售", 740, "沟通价格\n代客下单 合同"),
        ("系统管理员", 965, "超权限优惠审批\n按比例分摊"),
        ("客户", 1190, "线下付款\n上传付款凭证"),
        ("财务", 1415, "核对到账\n上传收款凭证"),
        ("客户", 1640, "按服务实例\n提交资料"),
        ("商务", 1865, "整体审核\n通过或驳回"),
        ("交付", 2080, "注册 补件 回签\n下证交付"),
    ]
    centers = []
    role_index = {role: i for i, role in enumerate(roles)}
    for role, x, label in nodes:
        y = lane_top + role_index[role] * lane_h + lane_h / 2
        box = (x - 92, y - 50, x + 92, y + 50)
        draw.rounded_rectangle(box, radius=18, fill="#" + ROLE_COLORS[role], outline="#7393B5", width=3)
        draw_centered(draw, box, label, node_font)
        centers.append((x + 92, y))
    for i in range(len(nodes) - 1):
        next_role, next_x, _ = nodes[i + 1]
        next_y = lane_top + role_index[next_role] * lane_h + lane_h / 2
        arrow(draw, centers[i], (next_x - 92, next_y))
    draw.text((285, 1090), "分支规则  无优惠订单按标准价付款  有优惠订单审批后更新应付金额  合同可在付款前后上传", font=font(27), fill="#455A70")
    image.save(path)


def make_role_flow(path: Path, role: str, steps: list[str], note: str):
    width, height = 2100, 380
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), f"{role}业务流程", font=font(42, True), fill="#102A43")
    node_w = min(270, int((width - 160 - (len(steps) - 1) * 70) / len(steps)))
    gap = int((width - 120 - len(steps) * node_w) / max(1, len(steps) - 1))
    y1, y2 = 120, 260
    centers = []
    for i, step in enumerate(steps):
        x1 = 60 + i * (node_w + gap)
        x2 = x1 + node_w
        box = (x1, y1, x2, y2)
        draw.rounded_rectangle(box, radius=20, fill="#" + ROLE_COLORS.get(role, LIGHT_BLUE), outline="#7393B5", width=3)
        draw_centered(draw, box, step, font(28))
        centers.append((x1, x2))
    for i in range(len(centers) - 1):
        arrow(draw, (centers[i][1] + 8, 190), (centers[i + 1][0] - 8, 190))
    draw.text((60, 315), note, font=font(25), fill="#52677D")
    image.save(path)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        if widths:
            cell.width = Inches(widths[idx])
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
            set_cell_border(cells[idx])
            set_cell_shading(cells[idx], "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[idx].paragraphs[0].paragraph_format.space_after = Pt(2)
            for run in cells[idx].paragraphs[0].runs:
                run.font.size = Pt(9)
        if widths:
            for idx, width in enumerate(widths):
                cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead:
        run = p.add_run(bold_lead)
        run.bold = True
    p.add_run(text)
    return p


def add_picture(doc: Document, path: Path, caption: str, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 103, 119)


def crop_preview(source: Path, destination: Path):
    image = Image.open(source).convert("RGB")
    max_height = int(image.width * 0.68)
    if image.height > max_height:
        image = image.crop((0, 0, image.width, max_height))
    image.thumbnail((1800, 1200), Image.Resampling.LANCZOS)
    image.save(destination, quality=92)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    styles["Subtitle"].font.size = Pt(12)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.bold = True

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(110)
    title.paragraph_format.space_after = Pt(18)
    title.add_run("EPR 下单系统业务流程与原型确认稿")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("散客选购 下单 付款 资料审核 注册办理 交付全流程")
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(28)
    info.add_run("版本 1.0\n确认日期 2026 年 10 月 12 日\n适用对象 产品 研发 测试 运营 财务 商务 交付")
    doc.add_page_break()

    add_heading(doc, "一 文档目的", 1)
    add_body(doc, "本确认稿汇总当前下单系统的端到端业务流程、角色边界、订单与服务状态、关键规则以及已实现原型页面。研发和业务团队可据此确认流程完整性，后续需求变更应同步更新流程、权限和验收场景。")
    add_body(doc, "系统采用订单与服务实例分离的设计。客户提交购物车后生成待付款订单；只有财务确认实际到账并上传收款凭证后，系统才生成服务实例。一个订单购买多份服务时，每份服务实例独立提交资料、审核、补件、注册和交付。", "核心结论  ")

    add_heading(doc, "二 角色与职责", 1)
    add_table(doc, ["角色", "主要职责", "关键权限", "不能执行"], [
        ["客户", "自助选购、下单、线下付款、上传凭证、提交资料、查看进度与交付文件", "购物车、订单、资料、补件、交付查看", "不能调整价格或确认到账"],
        ["销售", "维护客户、代客下单、沟通价格、设置打包优惠与合同", "客户管理、创建订单、打包优惠、合同维护", "不能确认收款或审核业务资料"],
        ["系统管理员", "审批超权限优惠、维护员工角色和系统配置", "优惠审批、权限、模板及全流程管理", "不代替财务进行到账确认"],
        ["财务", "核对线下到账、上传收款凭证、确认实付金额", "付款审核", "不能改优惠或审核业务资料"],
        ["商务", "按客户提交结构查看信息和材料，给出整体通过或驳回结论", "资料整体审核", "不逐字段逐文件下结论，不负责下证"],
        ["交付", "对接服务商、补件、二次回签、录入注册号和证书并完成服务", "任务分配、注册办理、交付完成", "不能确认收款或审批优惠"],
    ], widths=[0.9, 2.25, 1.75, 1.75])

    add_heading(doc, "三 端到端业务流程", 1)
    add_picture(doc, DIAGRAMS / "01_端到端业务流程.png", "图 1  EPR 下单系统端到端业务流程")
    add_body(doc, "标准价订单无需人工优惠，客户可按原价直接线下付款。需要优惠时，由销售为同一客户的多笔待付款订单设置打包价；系统按照各订单标准金额占比进行分摊，并保留标准价、优惠额、应付金额和优惠原因。")

    role_sections = [
        ("四 客户流程", "02_客户流程.png", "客户从选购到交付的主路径。资料提交发生在财务确认付款并生成服务实例之后。"),
        ("五 销售流程", "03_销售流程.png", "销售既可协助客户下单、设置同一客户的打包优惠，也可在付款前后维护合同；销售不能确认收款。"),
        ("六 系统管理员流程", "04_系统管理员流程.png", "优惠审批暂由系统管理员承担，审批结果必须形成价格快照和审计记录。"),
        ("七 财务流程", "05_财务流程.png", "客户上传付款凭证仅代表进入待确认状态；财务上传收款凭证并确认实付后，订单才变为已付款。"),
        ("八 商务流程", "06_商务流程.png", "商务按照客户提交时的分组结构查看全部字段和材料，但只给出服务实例级最终结论。驳回时必须填写可执行的补件建议。"),
        ("九 交付流程", "07_交付流程.png", "商务审核通过后进入注册办理。交付负责服务商递交、补件、二次回签、下证资料录入和最终完成。"),
    ]
    for heading, image_name, description in role_sections:
        doc.add_page_break()
        add_heading(doc, heading, 1)
        add_picture(doc, DIAGRAMS / image_name, f"图  {heading[0]}  {heading[2:]}")
        add_body(doc, description)

    doc.add_page_break()
    add_heading(doc, "十 核心业务规则", 1)
    add_table(doc, ["规则主题", "确认规则"], [
        ["订单生成", "客户购物车结算或销售代客下单后，订单默认状态为待付款。"],
        ["价格", "商品使用标准价；普通客户可直接按标准价付款。优惠是可选分支，不阻塞自助下单。"],
        ["打包优惠", "同一客户多笔待付款订单可设置打包价，系统按标准金额比例分摊优惠，尾差归最后一笔。"],
        ["合同", "销售可以在付款前或付款后上传合同，并记录合同编号、版本和审计轨迹。"],
        ["支付", "仅支持线下支付。客户上传付款凭证后进入付款待确认，不直接视为已付款。"],
        ["到账确认", "财务必须填写实际到账金额并上传收款凭证，确认后同步支付状态并生成服务实例。"],
        ["资料", "每个服务实例独立提交资料。历史企业和法人资料可复用，但客户可继续修改。"],
        ["商务审核", "按客户提交方式展示信息和材料，只给出整体通过或整体驳回结论；驳回必须写明建议。"],
        ["任务分配", "按国家和服务类型自动分配负责人，允许具备权限的人员人工调整。"],
        ["交付", "交付人员维护服务商递交、补件、回签、注册号、证书、账单和服务截止日期。"],
    ], widths=[1.35, 5.3])

    add_heading(doc, "十一 状态模型", 1)
    add_table(doc, ["对象", "状态链路", "说明"], [
        ["订单", "待付款 → 付款待确认 → 已付款", "付款凭证驳回时进入付款失败，客户可重新上传。"],
        ["服务实例", "草稿 → 已提交 → 审核中 → 注册中 → 已完成", "审核中可退回待补件，客户补件后重新进入审核中。"],
        ["优惠", "无优惠 或 待审批 → 已批准", "当前由系统管理员审批，订单保留标准价和优惠分摊结果。"],
        ["合同", "未上传 → 已上传 → 新版本", "合同与付款状态解耦，付款前后均可维护。"],
    ], widths=[1.1, 3.0, 2.55])

    doc.add_page_break()
    add_heading(doc, "十二 原型页面", 1)
    add_body(doc, "原型沿用当前 EPR FLOW 蓝白视觉体系。客户侧强调选购、订单、资料和进度；管理端根据登录角色动态展示菜单和操作按钮，后端同时执行权限校验。")
    preview_map = [
        ("01_客户登录页.png", "客户登录与注册入口"),
        ("02_客户工作台.png", "客户工作台与待办事项"),
        ("04_销售订单管理.png", "销售订单和合同管理"),
        ("05_财务收款审核.png", "财务付款凭证与实际到账审核"),
        ("06_商务资料审核.png", "商务待审核服务列表"),
        ("07_交付注册办理.png", "交付注册办理和下证信息"),
        ("08_系统管理员角色权限.png", "系统管理员员工角色与权限"),
    ]
    for index, (name, caption) in enumerate(preview_map, start=1):
        preview = ASSETS / f"preview_{index:02d}.jpg"
        crop_preview(SCREENSHOTS / name, preview)
        add_heading(doc, caption, 2)
        add_picture(doc, preview, f"原型页面 {index}  {caption}", width=6.45)

    doc.add_page_break()
    add_heading(doc, "十三 测试账号与场景", 1)
    add_table(doc, ["角色", "账号", "密码", "建议测试场景"], [
        ["客户", "demo@eprflow.com", "Demo123!", "自助下单、付款凭证、资料提交"],
        ["销售", "sales@eprflow.com", "Demo123!", "SHOW-SALES-001 代客下单和合同"],
        ["财务", "finance@eprflow.com", "Demo123!", "SHOW-FIN-001 到账确认"],
        ["商务", "business@eprflow.com", "Demo123!", "SHOWROLE20260001 至 003"],
        ["交付", "delivery@eprflow.com", "Demo123!", "SHOWROLE20260004 至 005"],
        ["系统管理员", "admin@eprflow.com", "Demo123!", "SHOW-ADMIN-001 和 002 优惠审批"],
    ], widths=[0.9, 2.0, 1.2, 2.55])

    add_heading(doc, "十四 验收确认清单", 1)
    for item in [
        "客户可按标准价自助下单，不要求必须申请优惠。",
        "销售可代客户创建订单并在付款前后上传合同。",
        "系统管理员可审批打包优惠，优惠金额能按订单标准价比例分摊。",
        "客户上传付款凭证后订单保持付款待确认，只有财务确认后才成为已付款。",
        "财务确认时必须上传收款凭证，订单显示实际到账金额。",
        "已付款订单按照购买数量生成独立服务实例。",
        "商务能按客户提交结构查看信息和材料，只提交整体通过或驳回结论。",
        "交付能处理注册、补件、回签和下证，完成后客户可以查看交付结果。",
        "不同角色只能看到相应菜单和操作，越权接口返回拒绝。",
        "关键价格、付款、审核、分配和交付动作均保留审计记录。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        p.add_run(item)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("EPR FLOW 业务流程与原型确认稿  2026-10-12")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(110, 120, 132)

    doc.save(DOCX_PATH)


def main():
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    make_overall_flow(DIAGRAMS / "01_端到端业务流程.png")
    make_role_flow(DIAGRAMS / "02_客户流程.png", "客户", ["选购服务", "购物车结算", "线下付款", "上传凭证", "提交资料", "补件", "接收交付"], "标准价可直接付款  优惠为可选分支  每个服务实例独立提交资料")
    make_role_flow(DIAGRAMS / "03_销售流程.png", "销售", ["建立客户", "代客下单", "沟通价格", "设置打包优惠", "上传合同", "跟进付款"], "合同可在付款前后上传  打包优惠按标准金额比例分摊")
    make_role_flow(DIAGRAMS / "04_系统管理员流程.png", "系统管理员", ["接收优惠申请", "校验同一客户", "审批打包价", "比例分摊", "记录审计", "维护角色模板"], "优惠审批暂由系统管理员负责  不代替财务确认到账")
    make_role_flow(DIAGRAMS / "05_财务流程.png", "财务", ["接收付款凭证", "核对银行到账", "核对付款主体", "填写实付金额", "上传收款凭证", "确认或驳回"], "确认到账后订单变为已付款并生成服务实例  驳回后客户可重新提交")
    make_role_flow(DIAGRAMS / "06_商务流程.png", "商务", ["接收已提交实例", "查看分组信息", "查看材料", "整体审核", "通过转交付", "驳回并给建议"], "不逐字段或逐文件给结论  驳回必须提供明确补件建议")
    make_role_flow(DIAGRAMS / "07_交付流程.png", "交付", ["接收任务", "提交服务商", "跟进补件", "二次回签", "录入注册号", "上传证书账单", "完成服务"], "按国家与服务类型自动分配负责人  允许有权限人员人工调整")
    build_docx()
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
